"""Regression tests for the code-review findings on the autopopulate batch.

Each test here was RED before its fix. They are grouped by the failure they
protect against, not by module, because several findings are the same failure
(an invisible skip, or a fabricated one) reached by different routes.
"""

import tempfile
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook

from autopopulate.models import Sidecar
from autopopulate.orchestrator import populate_phase


@pytest.fixture
def token_world(tmp_path):
    """A one-file templates_root + artifact dir whose payload has name AND industry."""
    root = tmp_path / "templates"
    root.mkdir()
    d = Document()
    d.add_paragraph().add_run("Prepared for [COMPANY]")
    d.save(root / "T.docx")

    art = tmp_path / "artifacts"
    (art / "output").mkdir(parents=True)
    (art / "output" / "company-overview.yaml").write_text(
        "exerciseId: company-overview\nconfirmedByClient: true\n"
        "payload:\n  company:\n    name: Acme Corp\n    industry: HVAC\n"
    )
    return root, art


# --- Finding 1: a mapped token absent from the template must be REPORTED ----

def test_token_missing_from_template_is_reported_even_when_another_token_matched(
    token_world, tmp_path
):
    root, art = token_world
    fm = tmp_path / "fm.yaml"
    fm.write_text(
        "company-overview:\n  - template: 'T.docx'\n    engine: docx-token\n"
        "    targets:\n"
        "      - source: company.name\n        token: '[COMPANY]'\n"
        "      - source: company.industry\n        token: '[INDUSTRY]'\n"
    )
    out = tmp_path / "out"
    result = populate_phase(art, "company-overview", out, root, fillmap_path=fm)

    assert any("[INDUSTRY]" in w for w in result.warnings), (
        "a mapped token that is absent from the template was dropped silently"
    )
    flag = out / "T.docx.FILL-MANUALLY.txt"
    assert flag.is_file() and "[INDUSTRY]" in flag.read_text(encoding="utf-8")
    # the token that DID match still landed
    assert Document(str(out / "T.docx")).paragraphs[0].text == "Prepared for Acme Corp"


def test_per_token_counts_are_reported_by_the_engine():
    from autopopulate.engines.docx import fill_docx_token_counts

    with tempfile.TemporaryDirectory() as td:
        p = Path(td) / "t.docx"
        d = Document()
        d.add_paragraph().add_run("[A] and [A] but no B")
        d.save(p)

        counts = fill_docx_token_counts(str(p), {"[A]": "x", "[B]": "y"})

    assert counts == {"[A]": 2, "[B]": 0}


# --- Finding 3: re-running the same exercise must not fabricate skips -------

def test_rerunning_the_same_exercise_does_not_fabricate_unfilled_reports(
    token_world, tmp_path
):
    root, art = token_world
    fm = tmp_path / "fm.yaml"
    fm.write_text(
        "company-overview:\n  - template: 'T.docx'\n    engine: docx-token\n"
        "    targets:\n      - source: company.name\n        token: '[COMPANY]'\n"
    )
    out = tmp_path / "out"
    populate_phase(art, "company-overview", out, root, fillmap_path=fm)
    second = populate_phase(art, "company-overview", out, root, fillmap_path=fm)

    assert Document(str(out / "T.docx")).paragraphs[0].text == "Prepared for Acme Corp"
    assert second.warnings == [], (
        "a correctly-filled file was reported as unfilled on re-run"
    )
    assert not (out / "T.docx.FILL-MANUALLY.txt").exists(), (
        "a correctly-filled file was flagged for manual filling"
    )


# --- Finding 2: a flag must not outlive the skips that caused it -----------

def test_flag_is_removed_once_its_skips_are_resolved(token_world, tmp_path):
    root, art = token_world
    thin = tmp_path / "thin.yaml"
    thin.write_text(
        "company-overview:\n  - template: 'T.docx'\n    engine: docx-token\n"
        "    targets:\n"
        "      - source: company.name\n        token: '[COMPANY]'\n"
        "      - source: company.tagline\n        token: '[TAGLINE]'\n"
    )
    out = tmp_path / "out"
    populate_phase(art, "company-overview", out, root, fillmap_path=thin)
    flag = out / "T.docx.FILL-MANUALLY.txt"
    assert flag.is_file(), "expected a flag naming the unfilled tagline"

    # Same exercise, now with the tagline mapping removed: nothing is unfilled.
    fixed = tmp_path / "fixed.yaml"
    fixed.write_text(
        "company-overview:\n  - template: 'T.docx'\n    engine: docx-token\n"
        "    targets:\n      - source: company.name\n        token: '[COMPANY]'\n"
    )
    out2 = tmp_path / "out2"
    populate_phase(art, "company-overview", out2, root, fillmap_path=thin)
    populate_phase(art, "company-overview", out2, root, fillmap_path=fixed)
    assert not (out2 / "T.docx.FILL-MANUALLY.txt").exists(), (
        "stale flag survived after its skips were resolved"
    )


# --- Finding 4: malformed YAML must surface as SidecarError/FillMapError ----

def test_malformed_sidecar_yaml_raises_sidecar_error(tmp_path):
    from autopopulate.sidecar import SidecarError, load_sidecar

    p = tmp_path / "bad.yaml"
    p.write_text("exerciseId: x\npayload: [unclosed\n")
    with pytest.raises(SidecarError):
        load_sidecar(p)


def test_malformed_fillmap_yaml_raises_fillmap_error(tmp_path):
    from autopopulate.fillmap import FillMapError, load_fillmap

    p = tmp_path / "bad.yaml"
    p.write_text("a: [unclosed\n")
    with pytest.raises(FillMapError):
        load_fillmap(p)


def test_one_bad_sidecar_does_not_cost_the_other_exercises_their_output(
    token_world, tmp_path
):
    from autopopulate.orchestrator import on_manifest_complete

    root, art = token_world
    (art / "output" / "positioning.yaml").write_text("payload: [unclosed\n")
    fm = tmp_path / "fm.yaml"
    fm.write_text(
        "company-overview:\n  - template: 'T.docx'\n    engine: docx-token\n"
        "    targets:\n      - source: company.name\n        token: '[COMPANY]'\n"
    )
    manifest = {
        "steps": [
            {"exerciseId": "positioning", "order": 1, "status": "complete"},
            {"exerciseId": "company-overview", "order": 2, "status": "complete"},
        ]
    }
    out = tmp_path / "out"
    results = on_manifest_complete(art, None, manifest, out, root, fillmap_path=fm)

    by_id = {r.exercise_id: r for r in results}
    assert any("positioning" in w for w in by_id["positioning"].warnings)
    assert by_id["company-overview"].filled == ["T.docx"], (
        "a malformed sidecar aborted the whole batch"
    )


# --- Finding 5: no engine may be accepted that cannot be executed ----------

def test_pptx_is_not_an_accepted_engine_while_unimplemented(tmp_path):
    from autopopulate.fillmap import ENGINES, FillMapError, load_fillmap

    assert "pptx" not in ENGINES
    p = tmp_path / "p.yaml"
    p.write_text(
        "a:\n  - template: 'D.pptx'\n    engine: pptx\n"
        "    targets:\n      - source: x\n        token: '[X]'\n"
    )
    with pytest.raises(FillMapError):
        load_fillmap(p)


def test_an_engine_crash_is_reported_not_fatal(token_world, tmp_path):
    """A template that is not a real docx must not kill the phase."""
    root, art = token_world
    (root / "Broken.docx").write_bytes(b"not a docx at all")
    fm = tmp_path / "fm.yaml"
    fm.write_text(
        "company-overview:\n  - template: 'Broken.docx'\n    engine: docx-token\n"
        "    targets:\n      - source: company.name\n        token: '[COMPANY]'\n"
    )
    out = tmp_path / "out"
    result = populate_phase(art, "company-overview", out, root, fillmap_path=fm)

    assert any("Broken.docx" in w for w in result.warnings)
    assert (out / "T.docx").is_file(), "the rest of the phase still had to be produced"


# --- Finding 6: numeric values must stay numeric in Excel ------------------

def test_numeric_payload_values_are_written_as_numbers_not_text(tmp_path):
    from openpyxl import Workbook

    root = tmp_path / "templates"
    root.mkdir()
    wb = Workbook()
    wb.active.title = "S"
    wb.save(root / "G.xlsx")

    art = tmp_path / "art"
    (art / "output").mkdir(parents=True)
    (art / "output" / "icp-prioritization.yaml").write_text(
        "exerciseId: icp-prioritization\nconfirmedByClient: true\n"
        "payload:\n  tiers:\n    - resourceAllocationPct: 60\n"
    )
    fm = tmp_path / "fm.yaml"
    fm.write_text(
        "icp-prioritization:\n  - template: 'G.xlsx'\n    engine: xlsx\n    sheet: 'S'\n"
        "    targets:\n      - source: tiers.0.resourceAllocationPct\n        cell: 'B2'\n"
    )
    out = tmp_path / "out"
    populate_phase(art, "icp-prioritization", out, root, fillmap_path=fm)

    cell = load_workbook(out / "G.xlsx")["S"]["B2"]
    assert cell.value == 60
    assert cell.data_type == "n", "number landed as text; SUM/sort/charts would ignore it"


def test_zero_and_false_are_real_values_not_empty(tmp_path):
    from autopopulate.fillmap import resolve_fills

    fm = tmp_path / "fm.yaml"
    fm.write_text(
        "a:\n  - template: 'T.xlsx'\n    engine: xlsx\n    sheet: 'S'\n"
        "    targets:\n"
        "      - source: count\n        cell: 'A1'\n"
        "      - source: flag\n        cell: 'A2'\n"
    )
    spec = resolve_fills("a", Sidecar("a", True, {"count": 0, "flag": False}),
                         fillmap_path=fm)[0]
    assert {t.dest: t.value for t in spec.targets} == {"A1": 0, "A2": False}
    assert spec.warnings == []


# --- Finding 7: a list of non-scalars must be a reported skip --------------

def test_join_refuses_non_scalar_list_elements(tmp_path):
    from autopopulate.fillmap import resolve_fills

    fm = tmp_path / "fm.yaml"
    fm.write_text(
        "a:\n  - template: 'T.docx'\n    engine: docx-token\n"
        "    targets:\n      - source: competitors\n        token: '[C]'\n        join: ', '\n"
    )
    payload = {"competitors": [{"name": "Globex Mechanical"}]}
    spec = resolve_fills("a", Sidecar("a", True, payload), fillmap_path=fm)[0]

    assert spec.targets == [], "a Python dict repr was written into a deliverable"
    assert len(spec.warnings) == 1


# --- Finding 8: block-level content controls must stay valid XML -----------

def test_block_level_dropdown_keeps_the_run_inside_a_paragraph(tmp_path):
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls, qn

    from autopopulate.engines.docx import set_docx_dropdown

    doc = Document()
    doc.element.body.append(parse_xml(
        f'<w:sdt {nsdecls("w")}><w:sdtPr><w:alias w:val="Block"/>'
        '<w:dropDownList><w:listItem w:displayText="Yes" w:value="Yes"/></w:dropDownList>'
        '</w:sdtPr><w:sdtContent><w:p><w:r><w:t>Choose.</w:t></w:r></w:p></w:sdtContent></w:sdt>'
    ))
    p = tmp_path / "block.docx"
    doc.save(p)

    assert set_docx_dropdown(str(p), "Block", "Yes") is True

    d = Document(str(p))
    for sdt in d.element.body.iter(qn("w:sdt")):
        content = sdt.find(qn("w:sdtContent"))
        kids = {k.tag.split("}")[1] for k in content}
        assert "r" not in kids, "a bare w:r as a direct child of w:sdtContent is invalid"
        texts = "".join(t.text or "" for t in content.iter(qn("w:t")))
        assert texts == "Yes"
        return
    raise AssertionError("sdt vanished")


# --- Finding 10: a malformed dotted index is a skip, not a crash -----------

def test_malformed_list_index_is_a_reported_skip_not_a_crash(tmp_path):
    from autopopulate.fillmap import resolve_fills

    fm = tmp_path / "fm.yaml"
    fm.write_text(
        "a:\n  - template: 'T.docx'\n    engine: docx-token\n"
        "    targets:\n      - source: tiers.--1.name\n        token: '[N]'\n"
    )
    spec = resolve_fills("a", Sidecar("a", True, {"tiers": [{"name": "x"}]}),
                         fillmap_path=fm)[0]
    assert spec.targets == []
    assert len(spec.warnings) == 1


# --- Finding 9: CLI must report a bad fill-map, not traceback --------------

def test_cli_reports_a_bad_fillmap_path(repo_root, mini_templates_root, tmp_path, capsys):
    from autopopulate.cli import main

    code = main([
        "--artifact-dir", str(repo_root / "tests/gtm_generator/fixtures/artifacts/happy-full"),
        "--exercise-id", "company-overview",
        "--output-dir", str(tmp_path / "o"),
        "--templates-root", str(mini_templates_root),
        "--fillmap", str(tmp_path / "does-not-exist.yaml"),
    ])
    assert code == 2
    assert "error:" in capsys.readouterr().err


# --- Portability: no machine- or account-specific paths in production ------

def test_production_code_has_no_hardcoded_absolute_or_account_paths():
    """The spec makes per-run output_dir + portability a hard requirement, and this
    repo is public — a baked-in home path or account id is both unportable and a leak.

    A WILDCARD account segment ("GoogleDrive-*") is fine: it names no person, works
    for any account, and is only a last-resort fallback behind the env var. What is
    banned is an absolute path, an email address, or a concrete account id.
    """
    import re

    banned = re.compile(
        r"""/Users/                      # absolute home path
          | [\w.+-]+@[\w-]+\.\w+         # email address
          | GoogleDrive-(?!\*)\S          # a CONCRETE Drive account
        """,
        re.VERBOSE,
    )
    pkg = Path(__file__).resolve().parents[2] / "autopopulate"
    offenders = [
        f"{f.relative_to(pkg.parent)}:{n}: {line.strip()}"
        for f in sorted(pkg.rglob("*.py"))
        for n, line in enumerate(f.read_text(encoding="utf-8").splitlines(), 1)
        if banned.search(line)
    ]
    assert offenders == [], (
        "hardcoded machine/account paths in production code:\n" + "\n".join(offenders)
    )


def test_templates_root_discovery_prefers_the_env_var(tmp_path, monkeypatch):
    from autopopulate.discovery import find_templates_root

    monkeypatch.setenv("STRATEGY_SPRINT_TEMPLATES", str(tmp_path))
    assert find_templates_root() == tmp_path


def test_templates_root_discovery_returns_none_when_nothing_is_found(tmp_path, monkeypatch):
    from autopopulate.discovery import find_templates_root

    monkeypatch.setenv("STRATEGY_SPRINT_TEMPLATES", str(tmp_path / "nope"))
    assert find_templates_root() is None
