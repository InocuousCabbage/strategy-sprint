"""World-asserting tests for the docx literal-token engine.

Every test reopens the file and asserts the value landed in the RIGHT paragraph /
run — not merely that the file changed.
"""

from docx import Document

from autopopulate.engines.docx import fill_docx_tokens


def test_company_token_replaced_in_body_and_table(make_docx_with_token):
    src = make_docx_with_token("[COMPANY]")
    n = fill_docx_tokens(src, {"[COMPANY]": "Acme"})

    d = Document(src)
    body_text = "\n".join(p.text for p in d.paragraphs)
    assert "[COMPANY]" not in body_text and "Acme" in body_text
    cell_text = d.tables[0].rows[0].cells[0].text
    assert "[COMPANY]" not in cell_text and "Acme" in cell_text
    assert n == 2


def test_value_lands_in_the_paragraph_that_held_the_token(make_docx_with_token):
    """Location specificity: the decoy paragraphs must be untouched."""
    src = make_docx_with_token("[COMPANY]")
    fill_docx_tokens(src, {"[COMPANY]": "Acme"})

    d = Document(src)
    assert d.paragraphs[0].text == "Annual Planning Worksheet"
    assert d.paragraphs[1].text == "Prepared for Acme — FY plan"
    assert d.paragraphs[2].text == "Do not edit this line."
    assert d.tables[0].rows[0].cells[1].text == "static"


def test_replacement_preserves_run_formatting(make_docx_with_token):
    """The whole point of run-level editing: bold/italic/size survive the fill."""
    src = make_docx_with_token("[COMPANY]")
    fill_docx_tokens(src, {"[COMPANY]": "Acme"})

    d = Document(src)
    p = d.paragraphs[1]
    assert len(p.runs) == 3, "runs were collapsed — formatting lost"
    target = p.runs[1]
    assert target.text == "Acme"
    assert target.bold is True
    assert target.italic is True
    assert target.font.size is not None
    assert p.runs[0].bold is not True  # neighbours keep their own formatting


def test_token_split_across_runs_is_replaced(make_docx_split_token):
    src = make_docx_split_token(("[COM", "PANY]"))
    n = fill_docx_tokens(src, {"[COMPANY]": "Acme"})

    d = Document(src)
    assert d.paragraphs[0].text == "Client: Acme"
    assert n == 1


def test_header_and_footer_tokens_replaced(make_docx_header_footer_token):
    src = make_docx_header_footer_token("[COMPANY]")
    n = fill_docx_tokens(src, {"[COMPANY]": "Acme"})

    d = Document(src)
    sec = d.sections[0]
    assert "Acme confidential" in "\n".join(p.text for p in sec.header.paragraphs)
    assert "(c) Acme" in "\n".join(p.text for p in sec.footer.paragraphs)
    assert n == 2


def test_unmapped_token_is_left_in_place_never_fabricated(make_docx_with_token):
    """No mapping -> the literal token stays visible for the consultant. Never invented."""
    src = make_docx_with_token("[COMPANY]")
    n = fill_docx_tokens(src, {"[SEGMENT]": "Core"})

    d = Document(src)
    assert "[COMPANY]" in d.paragraphs[1].text
    assert n == 0


def test_multiple_occurrences_in_one_run_all_replaced(tmp_path):
    from docx import Document as Doc

    doc = Doc()
    doc.add_paragraph().add_run("[COMPANY] and [COMPANY] again")
    p = tmp_path / "twice.docx"
    doc.save(p)

    n = fill_docx_tokens(str(p), {"[COMPANY]": "Acme"})

    d = Document(str(p))
    assert d.paragraphs[0].text == "Acme and Acme again"
    assert n == 2


def test_table_cell_is_filled_in_a_long_document(tmp_path):
    """Regression: paragraph de-duplication must not drop real paragraphs.

    Keying a 'seen' set on id(element) is unsound — lxml element proxies are
    transient, so ids get recycled and a later paragraph collides with a freed
    one's id and is silently skipped. A silently skipped fill is exactly the
    invisible-skip failure this component forbids.
    """
    from docx import Document as Doc

    doc = Doc()
    for i in range(60):
        doc.add_paragraph(f"filler {i}")
    doc.add_table(rows=1, cols=1).rows[0].cells[0].paragraphs[0].add_run("[COMPANY]")
    doc.add_paragraph().add_run("[COMPANY]")
    p = tmp_path / "long.docx"
    doc.save(p)

    n = fill_docx_tokens(str(p), {"[COMPANY]": "Acme"})

    d = Document(str(p))
    assert d.tables[0].rows[0].cells[0].text == "Acme"
    assert d.paragraphs[-1].text == "Acme"
    assert n == 2


def test_inherited_header_is_not_counted_twice(tmp_path):
    """A second section inheriting section 1's header must not double-count."""
    from docx import Document as Doc

    doc = Doc()
    doc.sections[0].header.paragraphs[0].add_run("[COMPANY] confidential")
    doc.add_paragraph("body")
    doc.add_section()  # inherits the header above
    assert doc.sections[1].header.is_linked_to_previous is True
    p = tmp_path / "twosec.docx"
    doc.save(p)

    n = fill_docx_tokens(str(p), {"[COMPANY]": "Acme"})

    d = Document(str(p))
    assert "Acme confidential" in d.sections[0].header.paragraphs[0].text
    assert n == 1


def test_paragraph_walk_yields_every_paragraph_exactly_once(tmp_path):
    """Deterministic guard on the walk itself.

    De-duplicating paragraphs by ``id(element)`` is unsound: lxml element proxies
    are transient, so ids get recycled and a later paragraph collides with a freed
    one's id and is dropped from the walk — a fill that silently never happens,
    which is exactly the invisible-skip failure this component forbids.
    """
    from collections import Counter

    from docx import Document as Doc

    from autopopulate.engines.docx import _all_paragraphs

    doc = Doc()
    expected = []
    for i in range(60):
        doc.add_paragraph(f"filler {i}")
        expected.append(f"filler {i}")
    cell = doc.add_table(rows=1, cols=1).rows[0].cells[0]
    cell.paragraphs[0].add_run("CELL")
    expected.append("CELL")
    nested = cell.add_table(rows=1, cols=1)
    nested.rows[0].cells[0].paragraphs[0].add_run("NESTED")
    expected.append("NESTED")
    doc.sections[0].header.paragraphs[0].add_run("HEADER")
    expected.append("HEADER")
    p = tmp_path / "walk.docx"
    doc.save(p)

    got = Counter(t for t in (x.text for x in _all_paragraphs(Document(str(p)))) if t)

    assert got == Counter(expected)
