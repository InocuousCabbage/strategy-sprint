"""World-asserting tests for docx structural fills: heading->blank paragraph, and
w:sdt dropdown content controls (the Positioning worksheet's shape).
"""

import pytest
from docx import Document

from autopopulate.engines.docx import (
    DropdownValueError,
    fill_docx_under_heading,
    set_docx_dropdown,
)
from tests.autopopulate.conftest import sdt_has_placeholder_flag, sdt_text

STATEMENT = "For regional property managers who cannot absorb downtime, Acme Corp is the HVAC partner that answers at 2am."


# --- heading -> blank paragraph -------------------------------------------

def test_fills_the_blank_paragraph_directly_below_the_heading(make_docx_with_headings):
    src = make_docx_with_headings()
    assert fill_docx_under_heading(src, "Positioning Statement", STATEMENT) is True

    d = Document(src)
    texts = [p.text for p in d.paragraphs]
    assert texts[0] == "", "decoy blank BEFORE the heading was overwritten"
    assert texts[1] == "Positioning Statement", "the heading itself was overwritten"
    assert texts[2] == STATEMENT, "value did not land in the blank below the heading"
    assert texts[3] == "Guidance: keep it to one sentence."
    assert texts[5] == "", "decoy blank under a DIFFERENT heading was overwritten"


def test_fill_uses_a_run_and_keeps_the_paragraph_style(make_docx_with_headings):
    src = make_docx_with_headings()
    before_style = Document(src).paragraphs[2].style.name

    fill_docx_under_heading(src, "Positioning Statement", STATEMENT)

    p = Document(src).paragraphs[2]
    assert len(p.runs) == 1 and p.runs[0].text == STATEMENT
    assert p.style.name == before_style


def test_unknown_heading_reports_not_filled(make_docx_with_headings):
    """VISIBLE SKIP: a heading we cannot find returns False so the caller records it."""
    src = make_docx_with_headings()
    assert fill_docx_under_heading(src, "No Such Heading", STATEMENT) is False
    assert [p.text for p in Document(src).paragraphs][2] == ""


def test_heading_with_no_blank_below_reports_not_filled(tmp_path):
    from docx import Document as Doc

    doc = Doc()
    doc.add_heading("Positioning Statement", level=1)
    doc.add_paragraph("already has content")
    p = tmp_path / "nogap.docx"
    doc.save(p)

    assert fill_docx_under_heading(str(p), "Positioning Statement", STATEMENT) is False
    assert Document(str(p)).paragraphs[1].text == "already has content"


# --- dropdown content controls --------------------------------------------

def test_sets_dropdown_to_an_allowed_value(make_docx_with_dropdown):
    src = make_docx_with_dropdown()
    assert set_docx_dropdown(src, "Product Type", "Vertical solution") is True
    assert sdt_text(src, "Product Type") == "Vertical solution"
    # the other control is untouched
    assert sdt_text(src, "Comparator") == "Choose an item."


def test_allowed_value_is_written_verbatim_including_trailing_space(make_docx_with_dropdown):
    """The template's list has "New way " with a trailing space; Word matches exactly."""
    src = make_docx_with_dropdown()
    set_docx_dropdown(src, "Product Type", "New way")  # caller trims; we write verbatim
    assert sdt_text(src, "Product Type") == "New way "


def test_rejects_a_value_not_in_the_controls_allowed_list(make_docx_with_dropdown):
    """Can-fail control: never invent an option Word would refuse."""
    src = make_docx_with_dropdown()
    with pytest.raises(DropdownValueError):
        set_docx_dropdown(src, "Product Type", "Cheapest option")
    assert sdt_text(src, "Product Type") == "Choose an item.", "document was mutated anyway"


def test_unknown_alias_reports_not_filled(make_docx_with_dropdown):
    """VISIBLE SKIP: no such control -> False, nothing written."""
    src = make_docx_with_dropdown()
    assert set_docx_dropdown(src, "Nonexistent Control", "Vertical solution") is False


def test_fill_clears_the_placeholder_flag(make_docx_with_dropdown):
    src = make_docx_with_dropdown()
    assert sdt_has_placeholder_flag(src, "Product Type") is True
    set_docx_dropdown(src, "Product Type", "Buy vs. Build")
    assert sdt_has_placeholder_flag(src, "Product Type") is False


def test_dropdown_fill_preserves_run_formatting(make_docx_with_dropdown):
    from docx.oxml.ns import qn

    src = make_docx_with_dropdown()
    set_docx_dropdown(src, "Comparator", "Horizontal solution not tailored to audience")

    d = Document(src)
    for sdt in d.element.body.iter(qn("w:sdt")):
        a = sdt.find(f"{qn('w:sdtPr')}/{qn('w:alias')}")
        if a is not None and a.get(qn("w:val")) == "Comparator":
            content = sdt.find(qn("w:sdtContent"))
            assert content.find(f"{qn('w:r')}/{qn('w:rPr')}/{qn('w:b')}") is not None, (
                "the run's bold formatting was discarded"
            )
            return
    raise AssertionError("Comparator sdt vanished")
