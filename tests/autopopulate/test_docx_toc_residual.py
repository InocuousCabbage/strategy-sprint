"""Finding 12 — and the CLASS it belongs to.

A token cached in a Table of Contents survived the fill: the TOC is a block-level
w:sdt, its entries repeat heading text verbatim, and the walk never entered
w:sdtContent. The residual was visible in the deliverable and reported nowhere.

The durable fix is not "also walk sdtContent" — that fixes today's container and
leaves tomorrow's (a text box, SmartArt, a comment, an embedded object) to fail the
same way. The durable fix is an OUTPUT SCAN: after a file is written, grep every
xml part of the finished package for any fill-map-declared token that still
remains, and report it.

Test priority mirrors that:
  PRIMARY  — the output scan reports a survivor in ANY container.
  SPECIFIC — the known TOC case is FILLED, not merely reported (a PDF export or a
             non-refreshing viewer would still show a literal otherwise).
  BELT     — w:updateFields asks Word to refresh too. Not load-bearing.
"""

import shutil
import zipfile

from docx import Document
from docx.oxml.ns import qn

from autopopulate.engines.docx import (
    fill_docx_token_counts,
    fill_docx_tokens,
    residual_docx_tokens,
)
from autopopulate.orchestrator import populate_phase


def _sdt_text(path):
    d = Document(path)
    return "\n".join(
        "".join(t.text or "" for t in c.iter(qn("w:t")))
        for sdt in d.element.body.iter(qn("w:sdt"))
        for c in [sdt.find(qn("w:sdtContent"))]
        if c is not None
    )


def _all_parts_literal_count(path, token):
    with zipfile.ZipFile(path) as z:
        return sum(
            z.read(n).decode("utf-8", "replace").count(token)
            for n in z.namelist()
            if n.lower().endswith(".xml")
        )


def _one_file_world(tmp_path, src, name):
    """A templates_root + confirmed artifact dir + fill-map around one docx."""
    root = tmp_path / "templates"
    root.mkdir(exist_ok=True)
    shutil.copy2(src, root / name)
    art = tmp_path / "art"
    (art / "output").mkdir(parents=True, exist_ok=True)
    (art / "output" / "company-overview.yaml").write_text(
        "exerciseId: company-overview\nconfirmedByClient: true\n"
        "payload:\n  company:\n    name: Acme Corp\n"
    )
    fm = tmp_path / "fm.yaml"
    fm.write_text(
        f"company-overview:\n  - template: '{name}'\n    engine: docx-token\n"
        "    targets:\n      - source: company.name\n        token: '[COMPANY]'\n"
    )
    return art, root, fm


# ===================== PRIMARY: the general output scan =====================

def test_output_scan_reports_a_survivor_in_a_container_the_fill_cannot_reach(
    make_docx_with_textbox_token, tmp_path
):
    """THE control. A text box is deliberately NOT the container we taught the walk
    about, so this tests the class: whatever hides a token, a survivor is reported.
    """
    art, root, fm = _one_file_world(
        tmp_path, make_docx_with_textbox_token("[COMPANY]"), "Box.docx"
    )
    out = tmp_path / "out"
    result = populate_phase(art, "company-overview", out, root, fillmap_path=fm)

    assert any("[COMPANY]" in w and "remain" in w.lower() for w in result.warnings), (
        f"a literal token survived in the output and was NOT reported: {result.warnings}"
    )
    flag = out / "Box.docx.FILL-MANUALLY.txt"
    assert flag.is_file(), "no .FILL-MANUALLY flag for a file with a surviving token"
    assert "[COMPANY]" in flag.read_text(encoding="utf-8")


def test_output_scan_greps_every_xml_part_not_the_paragraph_tree(
    make_docx_with_textbox_token,
):
    src = make_docx_with_textbox_token("[COMPANY]")
    fill_docx_tokens(src, {"[COMPANY]": "Acme Corp"})
    assert residual_docx_tokens(src, {"[COMPANY]": "Acme Corp"})["[COMPANY]"] == 1


def test_output_scan_catches_a_second_unrelated_container(
    make_docx_with_field_instruction,
):
    """A field instruction (w:instrText) is another container run.text never sees.
    One control, many containers — that is the point of scanning the package.
    """
    src = make_docx_with_field_instruction("[COMPANY]")
    fill_docx_tokens(src, {"[COMPANY]": "Acme Corp"})
    assert residual_docx_tokens(src, {"[COMPANY]": "Acme Corp"})["[COMPANY]"] == 1


def test_output_scan_is_clean_when_nothing_survives(make_docx_with_toc_cache):
    src = make_docx_with_toc_cache("[COMPANY]")
    fill_docx_tokens(src, {"[COMPANY]": "Acme Corp"})
    assert residual_docx_tokens(src, {"[COMPANY]": "Acme Corp"})["[COMPANY]"] == 0


def test_a_clean_fill_is_not_flagged(make_docx_with_toc_cache, tmp_path):
    art, root, fm = _one_file_world(
        tmp_path, make_docx_with_toc_cache("[COMPANY]"), "Annual.docx"
    )
    out = tmp_path / "out"
    result = populate_phase(art, "company-overview", out, root, fillmap_path=fm)

    assert result.warnings == []
    assert not (out / "Annual.docx.FILL-MANUALLY.txt").exists()
    assert result.filled == ["Annual.docx"]


# ============ SPECIFIC: the TOC case must be FILLED, not just reported ======

def test_token_cached_in_the_toc_is_filled(make_docx_with_toc_cache):
    """Reporting it is not enough: a PDF export or non-refreshing viewer would still
    show the literal. It has to be replaced in the file.
    """
    src = make_docx_with_toc_cache("[COMPANY]")
    counts = fill_docx_token_counts(src, {"[COMPANY]": "Acme Corp"})

    assert "[COMPANY]" not in _sdt_text(src), "TOC field cache still holds the token"
    assert "Acme Corp" in _sdt_text(src)
    assert counts["[COMPANY]"] == 2, "expected both the heading and its TOC cache"


def test_no_literal_token_survives_anywhere_after_a_toc_fill(make_docx_with_toc_cache):
    src = make_docx_with_toc_cache("[COMPANY]")
    fill_docx_tokens(src, {"[COMPANY]": "Acme Corp"})
    assert _all_parts_literal_count(src, "[COMPANY]") == 0


def test_widening_the_walk_did_not_lose_existing_coverage(make_docx_with_token):
    src = make_docx_with_token("[COMPANY]")
    assert fill_docx_token_counts(src, {"[COMPANY]": "Acme"})["[COMPANY]"] == 2


def test_inline_dropdown_controls_are_not_double_filled(make_docx_with_dropdown):
    """Inline w:sdt content is the dropdown engine's business. The token walk must
    not start editing it just because it now enters sdtContent.
    """
    src = make_docx_with_dropdown()
    assert fill_docx_token_counts(src, {"Choose an item.": "X"})["Choose an item."] == 0


# ===================== BELT: ask Word to refresh too ========================

def test_fill_marks_fields_for_update(make_docx_with_toc_cache):
    src = make_docx_with_toc_cache("[COMPANY]")
    fill_docx_tokens(src, {"[COMPANY]": "Acme Corp"})

    update = Document(src).settings.element.find(qn("w:updateFields"))
    assert update is not None and update.get(qn("w:val")) in ("true", "1")


def test_update_fields_is_not_duplicated_on_a_second_fill(make_docx_with_toc_cache):
    src = make_docx_with_toc_cache("[COMPANY]")
    fill_docx_tokens(src, {"[COMPANY]": "Acme Corp"})
    fill_docx_tokens(src, {"Acme Corp": "Acme Corporation"})

    assert len(Document(src).settings.element.findall(qn("w:updateFields"))) == 1


def test_output_scan_covers_parts_beyond_document_xml(tmp_path):
    """The scan must mean "every xml part", not "the body".

    A token in the document's core properties lives in docProps/core.xml — a part
    the fill engine never touches and the paragraph tree cannot see at all.
    """
    from docx import Document as Doc

    doc = Doc()
    doc.add_paragraph().add_run("body has no token")
    doc.core_properties.title = "[COMPANY] Annual Marketing Plan"
    p = tmp_path / "props.docx"
    doc.save(p)

    fill_docx_tokens(str(p), {"[COMPANY]": "Acme Corp"})

    assert residual_docx_tokens(str(p), {"[COMPANY]": "Acme Corp"})["[COMPANY]"] == 1
