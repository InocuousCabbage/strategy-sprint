"""The phase-done signal: sprint-manifest.yaml rows flipping to `complete`."""

import pytest

from autopopulate.manifest_watch import exercises_now_complete, load_manifest


def manifest(**statuses):
    return {
        "schemaVersion": "0.1.0",
        "steps": [
            {"exerciseId": eid, "order": i + 1, "status": st}
            for i, (eid, st) in enumerate(statuses.items())
        ],
    }


def test_returns_only_the_row_that_flipped_to_complete():
    prev = manifest(**{"company-overview": "complete", "icp-prioritization": "unrun"})
    cur = manifest(**{"company-overview": "complete", "icp-prioritization": "complete"})

    assert exercises_now_complete(prev, cur) == ["icp-prioritization"]


def test_a_row_complete_in_both_is_not_re_returned():
    """Idempotent: re-running the hook must not re-populate a done exercise."""
    m = manifest(**{"company-overview": "complete"})
    assert exercises_now_complete(m, m) == []


def test_several_flips_come_back_in_manifest_order():
    prev = manifest(
        **{"company-overview": "unrun", "icp-prioritization": "unrun", "positioning": "unrun"}
    )
    cur = manifest(
        **{"company-overview": "complete", "icp-prioritization": "unrun", "positioning": "complete"}
    )

    assert exercises_now_complete(prev, cur) == ["company-overview", "positioning"]


def test_no_previous_manifest_treats_every_complete_row_as_new():
    cur = manifest(**{"company-overview": "complete", "icp-prioritization": "unrun"})
    assert exercises_now_complete(None, cur) == ["company-overview"]


def test_a_row_that_regressed_and_completed_again_fires_again():
    prev = manifest(**{"positioning": "unrun"})
    cur = manifest(**{"positioning": "complete"})
    assert exercises_now_complete(prev, cur) == ["positioning"]


def test_a_row_that_went_backwards_does_not_fire():
    prev = manifest(**{"positioning": "complete"})
    cur = manifest(**{"positioning": "unrun"})
    assert exercises_now_complete(prev, cur) == []


def test_a_brand_new_row_arriving_complete_fires():
    prev = manifest(**{"company-overview": "complete"})
    cur = manifest(**{"company-overview": "complete", "positioning": "complete"})
    assert exercises_now_complete(prev, cur) == ["positioning"]


def test_reads_the_repos_real_manifest_fixture(repo_root):
    cur = load_manifest(
        repo_root / "tests/gtm_generator/fixtures/artifacts/happy-full/output/sprint-manifest.yaml"
    )
    assert exercises_now_complete(None, cur) == [
        "company-overview",
        "icp-prioritization",
        "positioning",
    ]


# --- the hook that actually populates --------------------------------------

def test_on_manifest_complete_populates_each_newly_complete_exercise(
    repo_root, mini_templates_root, tmp_path
):
    from docx import Document

    from autopopulate.orchestrator import on_manifest_complete

    artifact_dir = repo_root / "tests/gtm_generator/fixtures/artifacts/happy-full"
    cur = load_manifest(artifact_dir / "output" / "sprint-manifest.yaml")
    out = tmp_path / "out"

    results = on_manifest_complete(artifact_dir, None, cur, out, mini_templates_root)

    assert [r.exercise_id for r in results] == [
        "company-overview",
        "icp-prioritization",
        "positioning",
    ]
    annual = out / "3. Planning" / "3.3. Annual Planning" / "Annual Planning TEMPLATE.docx"
    assert Document(str(annual)).paragraphs[1].text == "Prepared for Acme Corp"
    pos = (
        out / "1. Foundations" / "1.2. Product Marketing Research"
        / "1.2.3. Product-Service" / "Positioning statement TEMPLATE.docx"
    )
    assert Document(str(pos)).paragraphs[1].text.startswith("For regional property managers")


def test_on_manifest_complete_reports_an_exercise_it_could_not_run(
    repo_root, mini_templates_root, tmp_path
):
    """A manifest row whose sidecar is missing must be reported, never swallowed."""
    from autopopulate.orchestrator import on_manifest_complete

    artifact_dir = repo_root / "tests/gtm_generator/fixtures/artifacts/happy-full"
    cur = {"steps": [{"exerciseId": "big-bets", "order": 8, "status": "complete"}]}
    out = tmp_path / "out"

    results = on_manifest_complete(artifact_dir, None, cur, out, mini_templates_root)

    assert len(results) == 1
    r = results[0]
    assert r.exercise_id == "big-bets"
    assert any("big-bets" in w for w in r.warnings)
    assert r.filled == []
