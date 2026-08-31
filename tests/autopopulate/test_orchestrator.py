"""Orchestrator tests.

Two rules dominate here and both are asserted on the FILESYSTEM, not on return
values alone:

- COPIES ONLY: the templates_root originals are byte-identical afterwards.
- VISIBLE SKIP: anything not filled is recorded — in the run's fill report AND in
  a .FILL-MANUALLY.txt beside the file. A skip that appears nowhere is the
  failure this component exists to prevent.
"""

import pytest
from docx import Document

from autopopulate.orchestrator import populate_phase
from autopopulate.sidecar import SidecarNotConfirmed
from tests.autopopulate.conftest import tree_hashes

ARTIFACTS = "tests/gtm_generator/fixtures/artifacts"


@pytest.fixture
def happy_artifact_dir(repo_root):
    return repo_root / ARTIFACTS / "happy-full"


def test_covered_file_is_mirrored_and_filled(happy_artifact_dir, mini_templates_root, tmp_path):
    out = tmp_path / "run-out"
    result = populate_phase(happy_artifact_dir, "company-overview", out, mini_templates_root)

    copy = out / "3. Planning" / "3.3. Annual Planning" / "Annual Planning TEMPLATE.docx"
    assert copy.is_file(), "covered template was not mirrored into output_dir"

    d = Document(str(copy))
    assert d.paragraphs[0].text == "Annual Marketing Plan"
    assert d.paragraphs[1].text == "Prepared for Acme Corp"
    assert d.paragraphs[1].runs[1].bold is True, "formatting lost in the copy"

    assert str(copy.relative_to(out)) in result.filled


def test_uncovered_file_is_copied_blank_and_flagged(
    happy_artifact_dir, mini_templates_root, tmp_path
):
    out = tmp_path / "run-out"
    result = populate_phase(happy_artifact_dir, "company-overview", out, mini_templates_root)

    copy = out / "2. Fuel + Engine" / "2.1. Fuel" / "Website Conversion Assessment TEMPLATE.xlsx"
    assert copy.is_file(), "uncovered template was not copied"
    flag = copy.with_suffix(copy.suffix + ".FILL-MANUALLY.txt")
    assert flag.is_file(), "uncovered file has no .FILL-MANUALLY flag"
    assert "no sprint source yet — fill manually" in flag.read_text(encoding="utf-8")
    assert str(copy.relative_to(out)) in result.copied_blank


def test_originals_are_byte_unchanged(happy_artifact_dir, mini_templates_root, tmp_path):
    before = tree_hashes(mini_templates_root)
    populate_phase(happy_artifact_dir, "company-overview", tmp_path / "o", mini_templates_root)
    assert tree_hashes(mini_templates_root) == before


def test_unconfirmed_sidecar_writes_nothing(repo_root, mini_templates_root, tmp_path):
    """Can-fail control: the gate runs before anything is copied."""
    out = tmp_path / "run-out"
    with pytest.raises(SidecarNotConfirmed):
        populate_phase(
            repo_root / ARTIFACTS / "unconfirmed", "company-overview", out, mini_templates_root
        )
    assert not out.exists() or list(out.rglob("*")) == []


def test_a_run_report_is_always_written(happy_artifact_dir, mini_templates_root, tmp_path):
    out = tmp_path / "run-out"
    result = populate_phase(happy_artifact_dir, "company-overview", out, mini_templates_root)

    report = out / "company-overview.fill-report.md"
    assert report.is_file()
    text = report.read_text(encoding="utf-8")
    assert "Annual Planning TEMPLATE.docx" in text
    assert "Website Conversion Assessment TEMPLATE.xlsx" in text
    assert str(report.relative_to(out)) == result.report_path


def test_unfilled_field_is_reported_and_flagged(
    happy_artifact_dir, mini_templates_root, tmp_path
):
    """THE non-negotiable: a field the sidecar cannot supply must surface twice —
    in the run report AND in a .FILL-MANUALLY flag beside the file it belongs to.
    """
    fm = tmp_path / "thin.yaml"
    fm.write_text(
        "company-overview:\n"
        "  - template: '3. Planning/3.3. Annual Planning/Annual Planning TEMPLATE.docx'\n"
        "    engine: docx-token\n"
        "    targets:\n"
        "      - source: company.name\n        token: '[COMPANY]'\n"
        "      - source: company.tagline\n        token: '[TAGLINE]'\n"
    )
    out = tmp_path / "run-out"
    result = populate_phase(
        happy_artifact_dir, "company-overview", out, mini_templates_root, fillmap_path=fm
    )

    assert any("company.tagline" in w for w in result.warnings)

    report = (out / "company-overview.fill-report.md").read_text(encoding="utf-8")
    assert "company.tagline" in report and "[TAGLINE]" in report

    copy = out / "3. Planning" / "3.3. Annual Planning" / "Annual Planning TEMPLATE.docx"
    flag = copy.with_suffix(copy.suffix + ".FILL-MANUALLY.txt")
    assert flag.is_file(), "a partially-filled file got no .FILL-MANUALLY flag"
    assert "company.tagline" in flag.read_text(encoding="utf-8")

    # the field that DID resolve still landed
    assert Document(str(copy)).paragraphs[1].text == "Prepared for Acme Corp"


def test_fillmap_naming_a_missing_template_is_reported(
    happy_artifact_dir, mini_templates_root, tmp_path
):
    """A renamed/moved template must not silently stop being filled."""
    fm = tmp_path / "ghost.yaml"
    fm.write_text(
        "company-overview:\n"
        "  - template: 'Nowhere/Ghost TEMPLATE.docx'\n"
        "    engine: docx-token\n"
        "    targets:\n      - source: company.name\n        token: '[COMPANY]'\n"
    )
    out = tmp_path / "run-out"
    result = populate_phase(
        happy_artifact_dir, "company-overview", out, mini_templates_root, fillmap_path=fm
    )

    assert any("Ghost TEMPLATE.docx" in w for w in result.warnings)
    assert "Ghost TEMPLATE.docx" in (out / "company-overview.fill-report.md").read_text(
        encoding="utf-8"
    )


def test_engine_rejection_is_reported_not_fatal(
    happy_artifact_dir, mini_templates_root, tmp_path
):
    """A heading that isn't in the document is a reportable skip, not a crash."""
    fm = tmp_path / "badheading.yaml"
    fm.write_text(
        "company-overview:\n"
        "  - template: '3. Planning/3.3. Annual Planning/Annual Planning TEMPLATE.docx'\n"
        "    engine: docx-heading\n"
        "    targets:\n"
        "      - source: company.name\n        heading: 'No Such Heading'\n"
    )
    out = tmp_path / "run-out"
    result = populate_phase(
        happy_artifact_dir, "company-overview", out, mini_templates_root, fillmap_path=fm
    )

    assert any("No Such Heading" in w for w in result.warnings)
    copy = out / "3. Planning" / "3.3. Annual Planning" / "Annual Planning TEMPLATE.docx"
    assert copy.is_file(), "the copy should still exist"
    flag = copy.with_suffix(copy.suffix + ".FILL-MANUALLY.txt")
    assert flag.is_file()


def test_folder_tree_is_mirrored_not_flattened(
    happy_artifact_dir, mini_templates_root, tmp_path
):
    out = tmp_path / "run-out"
    populate_phase(happy_artifact_dir, "company-overview", out, mini_templates_root)

    produced = {str(p.relative_to(out)) for p in out.rglob("*") if p.is_file()}
    assert "3. Planning/3.3. Annual Planning/Annual Planning TEMPLATE.docx" in produced
    assert "2. Fuel + Engine/2.1. Fuel/Website Conversion Assessment TEMPLATE.xlsx" in produced


def test_two_exercises_accumulate_into_one_output_dir(
    happy_artifact_dir, mini_templates_root, tmp_path
):
    """A phase runs several exercises into ONE output_dir. The second must not
    clobber the first by re-copying a pristine template over its filled copy.
    """
    out = tmp_path / "run-out"
    populate_phase(happy_artifact_dir, "company-overview", out, mini_templates_root)
    populate_phase(happy_artifact_dir, "positioning", out, mini_templates_root)

    annual = out / "3. Planning" / "3.3. Annual Planning" / "Annual Planning TEMPLATE.docx"
    assert Document(str(annual)).paragraphs[1].text == "Prepared for Acme Corp", (
        "the second exercise overwrote the first exercise's filled copy"
    )

    pos = (
        out / "1. Foundations" / "1.2. Product Marketing Research"
        / "1.2.3. Product-Service" / "Positioning statement TEMPLATE.docx"
    )
    assert Document(str(pos)).paragraphs[1].text.startswith("For regional property managers")


def test_a_fully_filled_file_keeps_no_flag(
    happy_artifact_dir, mini_templates_root, tmp_path
):
    """company-overview does not source the Positioning doc, so it flags it. When a
    later exercise fills it completely, the stale 'no sprint source' flag must go —
    a leftover flag sends the consultant to fill an already-filled file.
    """
    fm = tmp_path / "heading-only.yaml"
    fm.write_text(
        "positioning:\n"
        "  - template: '1. Foundations/1.2. Product Marketing Research/"
        "1.2.3. Product-Service/Positioning statement TEMPLATE.docx'\n"
        "    engine: docx-heading\n"
        "    targets:\n"
        "      - source: statement\n        heading: 'POSITIONING STATEMENT'\n"
    )
    out = tmp_path / "run-out"
    pos = (
        out / "1. Foundations" / "1.2. Product Marketing Research"
        / "1.2.3. Product-Service" / "Positioning statement TEMPLATE.docx"
    )
    flag = pos.with_suffix(pos.suffix + ".FILL-MANUALLY.txt")

    populate_phase(happy_artifact_dir, "company-overview", out, mini_templates_root)
    assert flag.is_file(), "expected a no-source flag from the first exercise"

    populate_phase(happy_artifact_dir, "positioning", out, mini_templates_root, fillmap_path=fm)
    assert not flag.exists(), "stale no-source flag survived a complete fill"


def test_a_partly_filled_file_swaps_the_flag_for_the_real_reasons(
    happy_artifact_dir, mini_templates_root, tmp_path
):
    """The shipped fill-map fills the Positioning statement but cannot yet source
    its two dropdowns. The flag must stop saying 'no sprint source yet' and start
    naming exactly which fields still need a human.
    """
    out = tmp_path / "run-out"
    pos = (
        out / "1. Foundations" / "1.2. Product Marketing Research"
        / "1.2.3. Product-Service" / "Positioning statement TEMPLATE.docx"
    )
    flag = pos.with_suffix(pos.suffix + ".FILL-MANUALLY.txt")

    populate_phase(happy_artifact_dir, "company-overview", out, mini_templates_root)
    assert "no sprint source yet" in flag.read_text(encoding="utf-8")

    populate_phase(happy_artifact_dir, "positioning", out, mini_templates_root)

    text = flag.read_text(encoding="utf-8")
    assert "no sprint source yet" not in text, "flag still claims the file has no source"
    assert "Product Type" in text and "Comparator" in text
    assert Document(str(pos)).paragraphs[1].text.startswith("For regional property managers")


def test_a_file_owned_by_an_earlier_exercise_is_reported_as_carried_over(
    happy_artifact_dir, mini_templates_root, tmp_path
):
    """Not re-touching a file must still be visible, not silent."""
    out = tmp_path / "run-out"
    populate_phase(happy_artifact_dir, "company-overview", out, mini_templates_root)
    result = populate_phase(happy_artifact_dir, "positioning", out, mini_templates_root)

    annual = "3. Planning/3.3. Annual Planning/Annual Planning TEMPLATE.docx"
    assert annual in result.carried_over
    assert annual not in result.copied_blank
    report = (out / "positioning.fill-report.md").read_text(encoding="utf-8")
    assert "carried over" in report.lower() and annual in report


def test_a_template_with_two_engines_gets_both_applied(
    happy_artifact_dir, mini_templates_root, tmp_path
):
    """The Positioning doc needs docx-heading AND docx-dropdown. Keying specs by
    template in a dict silently drops all but the last engine — the file looks
    processed and half its fills never happened.
    """
    fm = tmp_path / "two.yaml"
    tpl = "3. Planning/3.3. Annual Planning/Annual Planning TEMPLATE.docx"
    fm.write_text(
        "company-overview:\n"
        f"  - template: '{tpl}'\n"
        "    engine: docx-token\n"
        "    targets:\n      - source: company.name\n        token: '[COMPANY]'\n"
        f"  - template: '{tpl}'\n"
        "    engine: docx-heading\n"
        "    targets:\n"
        "      - source: company.industry\n        heading: 'OBJECTIVES'\n"
    )
    out = tmp_path / "run-out"
    populate_phase(
        happy_artifact_dir, "company-overview", out, mini_templates_root, fillmap_path=fm
    )

    d = Document(str(out / tpl))
    assert d.paragraphs[1].text == "Prepared for Acme Corp", "docx-token engine was dropped"
    assert "Commercial HVAC services" in "\n".join(p.text for p in d.paragraphs), (
        "docx-heading engine was dropped"
    )
