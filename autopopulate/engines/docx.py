"""DOCX fill engine.

Formatting discipline (non-negotiable): we edit ``run.text`` only. We NEVER assign
``paragraph.text = ...`` — that collapses every run in the paragraph into one and
throws away the bold/italic/size/style the consultant's template carries.

A token may be split across adjacent runs (Word does this constantly, e.g.
``[COM`` + ``PANY]``). We handle that by merging the minimal span of runs that
covers the token into the span's FIRST run, which keeps that run's formatting.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document


def _iter_paragraphs(container):
    """Yield every paragraph in a container, descending into tables (incl. nested)."""
    yield from getattr(container, "paragraphs", [])
    for table in getattr(container, "tables", []):
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_paragraphs(cell)


def _all_paragraphs(doc):
    """Every paragraph in the document, each exactly once.

    Deliberately does NOT de-duplicate by element identity: lxml element proxies
    are transient, so ``id()`` values get recycled and a later paragraph collides
    with a freed one's id and is dropped from the walk — a fill that silently
    never happens. Sections that inherit a header/footer are skipped instead, via
    ``is_linked_to_previous``, which is the actual document-level fact.
    """
    yield from _iter_paragraphs(doc)
    for section in doc.sections:
        for part in (section.header, section.footer):
            if part.is_linked_to_previous:
                continue  # inherited from an earlier section; already walked there
            yield from _iter_paragraphs(part)


def _replace_in_paragraph(paragraph, token: str, value: str) -> int:
    """Replace every ``token`` in one paragraph at run level. Returns the count."""
    count = 0
    search_from = 0
    while True:
        runs = paragraph.runs
        texts = [r.text for r in runs]
        joined = "".join(texts)
        idx = joined.find(token, search_from)
        if idx == -1:
            return count
        end = idx + len(token)

        starts = []
        pos = 0
        for t in texts:
            starts.append(pos)
            pos += len(t)

        first = last = None
        for i, t in enumerate(texts):
            s, e = starts[i], starts[i] + len(t)
            if first is None and s <= idx < e:
                first = i
            if first is not None and s < end <= e:
                last = i
                break
        if first is None or last is None:
            return count

        merged = "".join(texts[first : last + 1])
        local = idx - starts[first]
        runs[first].text = merged[:local] + value + merged[local + len(token) :]
        # Runs consumed by the merge are removed, not left as empty stubs.
        for r in runs[first + 1 : last + 1]:
            r._element.getparent().remove(r._element)

        count += 1
        # Step past the inserted value so a value containing the token can't loop.
        search_from = idx + len(value)


def fill_docx_token_counts(path: str | Path, tokens: dict[str, str]) -> dict[str, int]:
    """Replace each literal ``[TOKEN]`` throughout a docx COPY, preserving runs.

    Walks body paragraphs, table cells (nested included), and every section's
    header and footer. A token with no mapping is left in place verbatim — the
    consultant sees the unfilled placeholder rather than an invented value.

    Returns replacements PER TOKEN. The per-token breakdown is the point: a total
    would let a token that appears nowhere in the template hide behind a sibling
    token that matched, and that field would go unfilled with nothing reported.
    """
    p = Path(path)
    doc = Document(str(p))

    counts = {token: 0 for token in tokens}
    for paragraph in _all_paragraphs(doc):
        for token, value in tokens.items():
            counts[token] += _replace_in_paragraph(paragraph, token, str(value))

    if any(counts.values()):
        doc.save(str(p))
    return counts


def fill_docx_tokens(path: str | Path, tokens: dict[str, str]) -> int:
    """Total replacements made. See :func:`fill_docx_token_counts` for the detail."""
    return sum(fill_docx_token_counts(path, tokens).values())


class DropdownValueError(ValueError):
    """Raised when a value is not one of a dropdown content control's allowed options."""


_BODY_STYLE_PREFIXES = ("Heading", "Title", "Subtitle", "TOC")


def _is_body_paragraph(paragraph) -> bool:
    """True for a normal prose paragraph — not a heading/title/TOC entry."""
    name = (paragraph.style.name or "") if paragraph.style is not None else ""
    return not name.startswith(_BODY_STYLE_PREFIXES)


def fill_docx_under_heading(path: str | Path, heading_text: str, value: str) -> bool:
    """Write ``value`` into the first blank body paragraph AFTER ``heading_text``.

    Returns True if it filled, False if the heading was not found or had no blank
    paragraph beneath it before the next heading. False is a REPORTABLE skip: the
    caller records it so the consultant sees the field went unfilled — we never
    guess at a different location.
    """
    p = Path(path)
    doc = Document(str(p))
    paragraphs = doc.paragraphs

    for i, para in enumerate(paragraphs):
        if para.text.strip() != heading_text.strip():
            continue
        for candidate in paragraphs[i + 1 :]:
            if not _is_body_paragraph(candidate):
                break  # next heading reached — this section has no blank to fill
            if candidate.text.strip() == "":
                # Add a run rather than assigning .text, so the paragraph keeps its style.
                candidate.add_run(str(value))
                doc.save(str(p))
                return True
        return False
    return False


def _sdt_alias(sdt) -> str | None:
    from docx.oxml.ns import qn

    pr = sdt.find(qn("w:sdtPr"))
    if pr is None:
        return None
    alias = pr.find(qn("w:alias"))
    return None if alias is None else alias.get(qn("w:val"))


def _dropdown_items(sdt) -> list[str] | None:
    """The control's allowed options, read from the DOCUMENT — never hardcoded."""
    from docx.oxml.ns import qn

    pr = sdt.find(qn("w:sdtPr"))
    if pr is None:
        return None
    ddl = pr.find(qn("w:dropDownList"))
    if ddl is None:
        return None
    items = []
    for li in ddl.findall(qn("w:listItem")):
        items.append(li.get(qn("w:displayText")) or li.get(qn("w:value")) or "")
    return items


def set_docx_dropdown(path: str | Path, alias: str, value: str) -> bool:
    """Set the dropdown content control named ``alias`` to ``value``.

    ``value`` is matched against the control's own allowed list, compared trimmed,
    but the option is written back VERBATIM as the template spells it (the real
    Positioning template has "New way " with a trailing space, and Word only
    honours an exact match). A value outside the list raises ``DropdownValueError``
    rather than writing an option Word would reject.

    Returns True if set, False if no control carries that alias (a reportable skip).
    """
    from docx.oxml.ns import qn

    p = Path(path)
    doc = Document(str(p))

    for sdt in doc.element.body.iter(qn("w:sdt")):
        if _sdt_alias(sdt) != alias:
            continue

        allowed = _dropdown_items(sdt)
        if allowed is None:
            raise DropdownValueError(f"content control {alias!r} is not a dropdown list")

        match = next((o for o in allowed if o.strip() == str(value).strip()), None)
        if match is None:
            raise DropdownValueError(
                f"{value!r} is not an allowed option for {alias!r}; allowed: {allowed!r}"
            )

        content = sdt.find(qn("w:sdtContent"))
        # An INLINE sdt holds runs directly; a BLOCK-LEVEL one holds paragraphs.
        # A bare w:r is not a legal child of w:sdtContent, and Word reports a
        # document containing one as corrupt — so descend rather than append.
        runs = content.findall(qn("w:r")) or content.findall(f".//{qn('w:r')}")
        if not runs:
            parent = content.find(qn("w:p"))
            if parent is None:
                parent = content
            run = parent.makeelement(qn("w:r"), {})
            parent.append(run)
            runs = [run]
        # Keep the first run (and its rPr); write into its w:t. Drop the rest.
        first = runs[0]
        for extra in runs[1:]:
            extra.getparent().remove(extra)
        t = first.find(qn("w:t"))
        if t is None:
            t = first.makeelement(qn("w:t"), {})
            first.append(t)
        t.text = match
        t.set(
            "{http://www.w3.org/XML/1998/namespace}space", "preserve"
        )  # trailing space must survive

        # Clear the "showing placeholder" flag — the control now holds real content.
        pr = sdt.find(qn("w:sdtPr"))
        plc = pr.find(qn("w:showingPlcHdr")) if pr is not None else None
        if plc is not None:
            pr.remove(plc)

        doc.save(str(p))
        return True

    return False
