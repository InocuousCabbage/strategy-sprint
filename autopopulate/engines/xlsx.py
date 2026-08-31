"""XLSX fill engine.

Two rules the templates force on us:

1. **Data validation.** Many grid columns are dropdowns. Excel matches a dropdown
   option EXACTLY, and real templates ship options with trailing spaces ("Core ").
   So we match the caller's value trimmed, but write back the option as the SHEET
   spells it. A value outside the list is refused rather than written — a cell
   holding an option the dropdown rejects looks filled but is broken.

2. **Merged cells.** Only the top-left cell of a merge holds a value; a write to
   any other cell in the range is discarded by Excel (and openpyxl refuses it).
   We refuse it loudly instead of losing it quietly.

The workbook is loaded WITHOUT ``data_only`` so formulas survive the round-trip.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

from autopopulate.models import CellTarget

__all__ = [
    "CellTarget",
    "MergedCellError",
    "SheetNotFoundError",
    "ValidationValueError",
    "XlsxFillError",
    "fill_xlsx",
]


class XlsxFillError(Exception):
    """Base for xlsx fill problems."""


class SheetNotFoundError(XlsxFillError):
    """The workbook has no worksheet by that name."""


class ValidationValueError(XlsxFillError):
    """Value is not an option of the cell's list data-validation."""


class MergedCellError(XlsxFillError):
    """Target is a merged cell that is not its merge's top-left."""


def _inline_list_options(formula1: str) -> list[str] | None:
    """Options of an inline list validation, e.g. ``"Core ,Scaling"`` -> [...]."""
    f = (formula1 or "").strip()
    if not (f.startswith('"') and f.endswith('"')):
        return None
    return f[1:-1].split(",")


def _range_list_options(wb, formula1: str) -> list[str] | None:
    """Options of a range-backed list validation, e.g. ``Lists!$A$1:$A$2``."""
    f = (formula1 or "").strip().lstrip("=")
    if "!" not in f:
        return None
    sheet_name, ref = f.rsplit("!", 1)
    sheet_name = sheet_name.strip("'")
    if sheet_name not in wb.sheetnames:
        return None
    src = wb[sheet_name]
    try:
        min_col, min_row, max_col, max_row = range_boundaries(ref.replace("$", ""))
    except (ValueError, TypeError):
        return None
    options = []
    for row in src.iter_rows(
        min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col
    ):
        for cell in row:
            if cell.value is not None:
                options.append(str(cell.value))
    return options or None


def _validation_options(wb, ws, coord: str) -> list[str] | None:
    """Allowed options for ``coord``, or None if it carries no list validation."""
    for dv in ws.data_validations.dataValidation:
        if dv.type != "list" or coord not in dv.sqref:
            continue
        return _inline_list_options(dv.formula1) or _range_list_options(wb, dv.formula1)
    return None


def _merge_conflict(ws, coord: str) -> str | None:
    """The merge range whose non-top-left cell ``coord`` is, if any."""
    cell = ws[coord]
    for rng in ws.merged_cells.ranges:
        if (rng.min_row <= cell.row <= rng.max_row) and (
            rng.min_col <= cell.column <= rng.max_col
        ):
            if (cell.row, cell.column) != (rng.min_row, rng.min_col):
                return str(rng)
    return None


def fill_xlsx(path: str | Path, sheet: str, cell_targets: list[CellTarget]) -> int:
    """Write ``cell_targets`` into ``sheet`` of an xlsx COPY. Returns the count.

    Every target is validated BEFORE anything is written, so a rejected target
    never leaves a half-filled sheet behind.
    """
    p = Path(path)
    wb = load_workbook(str(p))  # NOT data_only — formulas must survive
    if sheet not in wb.sheetnames:
        raise SheetNotFoundError(f"{p} has no sheet {sheet!r}; has {wb.sheetnames!r}")
    ws = wb[sheet]

    resolved: list[tuple[str, object]] = []
    for target in cell_targets:
        coord, value = target.coord, target.value

        conflict = _merge_conflict(ws, coord)
        if conflict is not None:
            raise MergedCellError(
                f"{coord} is not the top-left of merge {conflict} in {sheet!r}; "
                f"Excel would discard the write"
            )

        options = _validation_options(wb, ws, coord)
        if options is not None:
            match = next((o for o in options if o.strip() == str(value).strip()), None)
            if match is None:
                raise ValidationValueError(
                    f"{value!r} is not an option of the list validation on "
                    f"{sheet}!{coord}; allowed: {options!r}"
                )
            value = match  # write the sheet's own spelling, trailing space and all

        resolved.append((coord, value))

    for coord, value in resolved:
        ws[coord] = value

    if resolved:
        wb.save(str(p))
    return len(resolved)
